#!/usr/bin/env python3
"""Builds pitstop-deck.pptx from the screenshots produced by
capture-slides.js plus the deck's own SCRIPT object (speaker notes), so the
notes never drift from the on-screen `S` script.

Usage:
    pip3 install python-pptx
    node capture-slides.js /tmp/pitstop-pptx-build
    python3 build-pptx.py /tmp/pitstop-pptx-build
"""
import re
import sys
import os
from pptx import Presentation
from pptx.util import Inches
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
HTML = os.path.join(HERE, '..', 'pitstop-deck.html')
OUT = os.path.join(HERE, '..', 'pitstop-deck.pptx')
IMG_DIR = sys.argv[1] if len(sys.argv) > 1 else '/tmp/pitstop-pptx-build'

with open(HTML, 'r', encoding='utf-8') as f:
    html = f.read()

m = re.search(r'var SCRIPT\s*=\s*\{(.*?)\n  \};', html, re.S)
body = m.group(1)
entries = re.findall(r'(\d+):\[(.*?)\]\s*,?\s*(?=\n\s*\d+:\[|\Z)', body, re.S)

def parse_strings(s):
    # two JS single-quoted strings, escaped quotes as \'
    out, cur, i, in_str, esc = [], '', 0, False, False
    while i < len(s):
        c = s[i]
        if in_str:
            if esc:
                cur += c
                esc = False
            elif c == '\\':
                esc = True
            elif c == "'":
                in_str = False
                out.append(cur)
                cur = ''
            else:
                cur += c
        else:
            if c == "'":
                in_str = True
        i += 1
    return out

SCRIPT = {}
for num, raw in entries:
    strs = parse_strings(raw)
    if len(strs) >= 2:
        SCRIPT[int(num)] = (strs[0], strs[1])

MAIN_SLIDES = 20

def clean(s):
    return (s.replace('&rsquo;', "'").replace('&ldquo;', '"').replace('&rdquo;', '"')
             .replace('&mdash;', '\u2014').replace('&amp;', '&').replace('&nbsp;', ' '))

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

P_NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'

def add_transition(slide):
    """Every slide-to-slide move in the deck is the same fast cross-fade with
    a small push (.slide{opacity 0->1, translateX 26px->0}, 0.34s ease).
    'fade' + 'fast' speed is pptx's closest native equivalent."""
    xml = (f'<p:transition xmlns:p="{P_NS}" spd="fast" '
           'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" p14:dur="340">'
           '<p:fade/></p:transition>')
    slide._element.append(etree.fromstring(xml))

for n in range(1, MAIN_SLIDES + 1):
    slide = prs.slides.add_slide(blank)
    img = os.path.join(IMG_DIR, f'slide-{n:02d}.png')
    slide.shapes.add_picture(img, 0, 0, width=prs.slide_width, height=prs.slide_height)
    cue, text = SCRIPT.get(n, ('', ''))
    tf = slide.notes_slide.notes_text_frame
    tf.text = clean(cue)
    p2 = tf.add_paragraph()
    p2.text = clean(text)
    add_transition(slide)

prs.save(OUT)
print('wrote', OUT)
